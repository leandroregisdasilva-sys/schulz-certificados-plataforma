from __future__ import annotations

import io
import re
import sys
import zipfile
from datetime import date, datetime, time
from pathlib import Path
from typing import Iterable, Sequence

from openpyxl import Workbook

RAIZ = Path(__file__).resolve().parent.parent


def tabela_simples(pontos: Sequence[tuple], *, rotulos_leitura: Sequence[str] | None = None,
                   cabecalho: Sequence[str] | None = None, meta: dict[str, object] | None = None,
                   aba: str = "Leituras", linhas_topo: int = 0) -> bytes:
    n_leit = max(len(p) - 3 for p in pontos) if pontos else 2
    rotulos = list(rotulos_leitura or [f"Leitura {i}" for i in range(1, n_leit + 1)])
    cab = list(cabecalho or ["Ponto", "Valor padrão (mm)", "LEA (mm)", *rotulos])
    wb = Workbook()
    ws = wb.active
    ws.title = aba
    linha = 1
    for chave, valor in (meta or {}).items():
        ws.cell(row=linha, column=1, value=chave)
        ws.cell(row=linha, column=2, value=valor)
        linha += 1
    linha += linhas_topo
    for c, t in enumerate(cab, start=1):
        ws.cell(row=linha, column=c, value=t)
    for p in pontos:
        linha += 1
        for c, v in enumerate(p, start=1):
            ws.cell(row=linha, column=c, value=v)
    return _bytes(wb)


def tabela_tolerancias(pontos: Sequence[tuple], *, meta: dict[str, object] | None = None) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Registro"
    linha = 1
    for chave, valor in (meta or {}).items():
        ws.cell(row=linha, column=2, value=chave)
        ws.cell(row=linha, column=3, value=valor)
        linha += 2
    linha += 1
    cab = ["Parâmetro avaliado", "Valor nominal", "Tol. +", "Tol. -", "Medida 1", "Medida 2"]
    for c, t in enumerate(cab, start=2):
        ws.cell(row=linha, column=c, value=t)
    for p in pontos:
        linha += 1
        for c, v in enumerate(p, start=2):
            ws.cell(row=linha, column=c, value=v)
    return _bytes(wb)


def calypso(blocos: Sequence[Sequence[tuple]], *, plano: str = "999.0001_EXEMPLO",
            metrologista: str = "Metrologista Exemplo", mmc: int = 500011,
            fornecedor: str = "Fornecedor X", data: datetime | None = None,
            rodape: bool = True) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Report"
    ws["B2"] = "Calypso Measuring Result"
    for col, rot in (("B", "Measurement Plan"), ("D", "Date"), ("F", "Order")):
        ws[f"{col}4"] = rot
    ws["B5"] = plano
    ws["D5"] = data or datetime(2026, 9, 1)
    for col, rot in (("B", "Drawing No."), ("D", "Time"), ("F", "Part No.")):
        ws[f"{col}7"] = rot
    ws["D8"] = time(9, 30, 0)
    ws["F8"] = 1
    for col, rot in (("B", "Metrologista"), ("D", "MMC"), ("F", "Fornecedor")):
        ws[f"{col}10"] = rot
    ws["B11"], ws["D11"], ws["F11"] = metrologista, mmc, fornecedor
    for c, t in enumerate(["Characteristic", "Actual", "Nominal", "Upper Tol", "Lower Tol", "Deviation"], start=1):
        ws.cell(row=13, column=c, value=t)
    linha = 13
    for bloco in blocos:
        for carac, actual, nominal, up, lo in bloco:
            linha += 1
            for c, v in enumerate([carac, actual, nominal, up, lo, actual - nominal], start=1):
                ws.cell(row=linha, column=c, value=v)
        linha += 1
    if rodape:
        for c, v in enumerate(["Nota: incerteza do equipamento conforme certificado", None, None], start=1):
            ws.cell(row=linha + 1, column=c, value=v)
        for k, v in enumerate(["2", "0", "1"]):
            ws.cell(row=linha + 2 + k, column=1, value=v)
    return _bytes(wb)


def registro_schulz(pontos: Sequence[tuple], *, cliente="Setor Exemplo", tecnico="Técnico Exemplo",
                    codigo="000.0000", denominacao="Dispositivo de controle",
                    fabricante="Não consta", data=datetime(2026, 9, 1, 10, 0, 0)) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Registro"
    ws["C1"] = "Registro - Calibração de Dispositivo de controle"
    ws["A3"] = "Dados de registro"
    ws["B5"], ws["C5"] = "Data:", data
    ws["G5"], ws["H5"] = "Temperatura ambiente:", "20,2°"
    ws["B7"], ws["C7"], ws["G7"] = "Cliente solicitante:", cliente, "OS:"
    ws["B9"], ws["C9"] = "Técnico executante:", tecnico
    ws["B12"], ws["C12"] = "Padrão utilizado:", 1
    ws["A16"] = "Dados do instrumento"
    ws["B18"], ws["C18"] = "Código:", codigo
    ws["B20"], ws["C20"] = "Denominação:", denominacao
    ws["B22"], ws["C22"] = "Valor nominal:", "Exclusivo"
    ws["B24"], ws["C24"] = "Fabricante:", fabricante
    ws["A27"] = "Resultados obtidos (mm)"
    cab = ["Parâmetro avaliado", "Valor nominal", "Tol. +", "Tol. -", "Medida 1", "Medida 2",
           "Média", "Desvio padrão", None, "ABS do valor nominal", "tol+", "tol-"]
    for c, t in enumerate(cab, start=2):
        if t:
            ws.cell(row=29, column=c, value=t)
    linha = 29
    for nome, nominal, tp, tn, m1, m2 in pontos:
        linha += 1
        media = (m1 + m2) / 2
        dp = abs(m1 - m2) / (2 ** 0.5)
        for c, v in enumerate([nome, nominal, tp, tn, m1, m2, media, dp, None, abs(nominal), "Rep", "Rep"], start=2):
            ws.cell(row=linha, column=c, value=v)
    for _ in range(10):
        linha += 1
        for c in range(2, 11):
            ws.cell(row=linha, column=c, value=" ")
        ws.cell(row=linha, column=12, value="#VALUE!")
        ws.cell(row=linha, column=13, value="#VALUE!")
    return _bytes(wb)


def _bytes(wb: Workbook) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


_TRANSICIONAL_PARA_ESTRITO = (
    ("http://schemas.openxmlformats.org/spreadsheetml/2006/main", "http://purl.oclc.org/ooxml/spreadsheetml/main"),
    ("http://schemas.openxmlformats.org/officeDocument/2006/relationships", "http://purl.oclc.org/ooxml/officeDocument/relationships"),
    ("http://schemas.openxmlformats.org/drawingml/2006/main", "http://purl.oclc.org/ooxml/drawingml/main"),
)


def para_estrito(dados: bytes) -> bytes:
    saida = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(dados)) as zi, zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as zo:
        for item in zi.infolist():
            conteudo = zi.read(item.filename)
            if item.filename.endswith((".xml", ".rels")):
                texto = conteudo.decode("utf-8")
                for transicional, estrito in _TRANSICIONAL_PARA_ESTRITO:
                    texto = texto.replace(transicional, estrito)
                if item.filename == "xl/workbook.xml":
                    texto = re.sub(r"<workbook ", '<workbook conformance="strict" ', texto, count=1)
                conteudo = texto.encode("utf-8")
            zo.writestr(item, conteudo)
    return saida.getvalue()


PONTOS_APROVADO = [
    ("DIAMETRO_A", 10.000, 0.010, 10.0021, 10.0027, 10.0024),
    ("DIAMETRO_B", 25.000, 0.010, 24.9962, 24.9958, 24.9960),
    ("DIAMETRO_C", 40.000, 0.010, 40.0050, 40.0047, 40.0049),
    ("COMPRIMENTO_A", 75.000, 0.020, 75.0110, 75.0096, 75.0103),
    ("COMPRIMENTO_B", 100.000, 0.020, 99.9915, 99.9921, 99.9918),
    ("PLANICIDADE", 0.000, 0.050, 0.0123, 0.0131, 0.0127),
    ("PERPENDICULARISMO", 0.000, 0.030, 0.0088, 0.0094, 0.0091),
    ("PARALELISMO", 0.000, 0.030, 0.0150, 0.0162, 0.0156),
    ("LIMITE_EXATO", 1.000, 0.010, 1.0100, 1.0100, 1.0100),
    ("RUGOSIDADE_REF", 5.000, 0.005, 4.9981, 4.9979, 4.9980),
]
PONTOS_REPROVADO = [
    ("DIAMETRO_A", 10.000, 0.010, 10.0021, 10.0027, 10.0024),
    ("DIAMETRO_B", 25.000, 0.010, 25.0152, 25.0148, 25.0150),
    ("DIAMETRO_C", 40.000, 0.010, 39.9890, 39.9894, 39.9892),
    ("COMPRIMENTO_A", 75.000, 0.020, 75.0110, 75.0096, 75.0103),
    ("COMPRIMENTO_B", 100.000, 0.020, 99.9915, 99.9921, 99.9918),
    ("PLANICIDADE", 0.000, 0.050, 0.0123, 0.0131, 0.0127),
    ("PERPENDICULARISMO", 0.000, 0.030, 0.0088, 0.0094, 0.0091),
    ("PARALELISMO", 0.000, 0.030, 0.0150, 0.0162, 0.0156),
]
BLOCO_CALYPSO_1 = [
    ("PLANICIDADE_REF", 0.002100, 0.0, 0.05, 0.0),
    ("ESFERA_DIAMETRO 1", 12.0034, 12.0, 0.01, -0.01),
    ("ESFERA_POSICAO 1", 0.0412, 0.0, 0.10, 0.0),
    ("FURO_DIAMETRO 2", 8.0061, 8.0, 0.01, -0.01),
    ("FURO_POSICAO 2", 0.0733, 0.0, 0.10, 0.0),
    ("ENTRE_CENTROS", 0.0215, 0.0, 0.04, 0.0),
]
BLOCO_CALYPSO_2 = [
    ("PLANICIDADE_REF", 0.002300, 0.0, 0.05, 0.0),
    ("ESFERA_DIAMETRO 1", 12.0030, 12.0, 0.01, -0.01),
    ("ESFERA_POSICAO 1", 0.0398, 0.0, 0.10, 0.0),
    ("FURO_DIAMETRO 2", 8.0057, 8.0, 0.01, -0.01),
    ("FURO_POSICAO 2", 0.0741, 0.0, 0.10, 0.0),
    ("ENTRE_CENTROS", 0.0221, 0.0, 0.04, 0.0),
]
PONTOS_REGISTRO = [
    ("PLANICIDADE_REF", 0.0, 0.05, 0.0, 0.0021, 0.0023),
    ("DIAMETRO 1", 6.075, 0.01, -0.01, 6.0781, 6.0779),
    ("POSICAO 1", 0.0, 0.1, 0.0, 0.0612, 0.0587),
    ("DIAMETRO 2", 16.964, 0.01, -0.01, 16.9668, 16.9671),
    ("POSICAO 2", 0.0, 0.1, 0.0, 0.0420, 0.0433),
]


CALYPSO_APROVADO = [
    ("PLANICIDADE_REF", 0.0021, 0.0023, 0.0, 0.05, 0.0),
    ("ENTRE CENTROS_REF", 0.0268, 0.0272, 0.0, 0.04, 0.0),
    ("62P_DIÂMETRO 1", 6.0791, 6.0787, 6.075, 0.01, -0.01),
    ("62P_POSIÇÃO 1", 0.0430, 0.0418, 0.0, 0.10, 0.0),
    ("62P_DIÂMETRO 2", 6.0779, 6.0783, 6.075, 0.01, -0.01),
    ("62P_POSIÇÃO 2", 0.0245, 0.0251, 0.0, 0.10, 0.0),
    ("13P_DIÂMETRO 1", 16.9681, 16.9677, 16.964, 0.01, -0.01),
    ("13P_POSIÇÃO 1", 0.0559, 0.0547, 0.0, 0.10, 0.0),
    ("13P_DIÂMETRO 2", 16.9702, 16.9698, 16.964, 0.01, -0.01),
    ("13P_POSIÇÃO 2", 0.0236, 0.0240, 0.0, 0.10, 0.0),
    ("106P_DIÂMETRO 1", 6.3461, 6.3459, 6.345, 0.01, -0.01),
    ("106P_POSIÇÃO 1", 0.0132, 0.0128, 0.0, 0.10, 0.0),
    ("106P_DIÂMETRO 2", 6.3489, 6.3491, 6.345, 0.01, -0.01),
    ("106P_POSIÇÃO 2", 0.0868, 0.0862, 0.0, 0.10, 0.0),
]
FORA_DA_TOLERANCIA = {"13P_POSIÇÃO 1": (0.1259, 0.1247), "62P_DIÂMETRO 2": (6.0862, 6.0858)}


def blocos_calypso(pontos, alteracoes=None):
    alteracoes = alteracoes or {}
    bloco1, bloco2 = [], []
    for nome, m1, m2, nominal, tp, tn in pontos:
        m1, m2 = alteracoes.get(nome, (m1, m2))
        bloco1.append((nome, m1, nominal, tp, tn))
        bloco2.append((nome, m2, nominal, tp, tn))
    return [bloco1, bloco2]


PONTOS_REGISTRO_NOVO = [
    ("PLANICIDADE_BASE", 0.0, 0.03, 0.0, 0.0112, 0.0118),
    ("DIAMETRO GUIA 1", 12.0, 0.01, -0.01, 12.0031, 12.0035),
    ("DIAMETRO GUIA 2", 12.0, 0.01, -0.01, 11.9978, 11.9974),
    ("POSICAO GUIA 1", 0.0, 0.05, 0.0, 0.0215, 0.0221),
    ("COMPRIMENTO TOTAL", 85.0, 0.02, -0.02, 85.0071, 85.0065),
]


def gerar_todos(destino: Path) -> list[Path]:
    destino.mkdir(parents=True, exist_ok=True)
    data = datetime(2026, 10, 20, 9, 30)
    arquivos = {
        "01_calypso_703-1208_aprovado.xlsx": calypso(blocos_calypso(CALYPSO_APROVADO), plano="703.1208_DEMO",
                                                     metrologista="Técnico Demonstração", mmc=500011,
                                                     fornecedor="Não consta", data=data),
        "02_calypso_703-1208_reprovado.xlsx": calypso(blocos_calypso(CALYPSO_APROVADO, FORA_DA_TOLERANCIA),
                                                      plano="703.1208_DEMO", metrologista="Técnico Demonstração",
                                                      mmc=500011, fornecedor="Não consta", data=data),
        "03_registro_instrumento_novo_strict.xlsx": para_estrito(registro_schulz(
            PONTOS_REGISTRO_NOVO, cliente="Montagem Compressores", tecnico="Técnico Demonstração",
            codigo="900.0001", denominacao="Dispositivo de controle", fabricante="Não consta", data=data)),
    }
    caminhos = []
    for nome, conteudo in arquivos.items():
        caminho = destino / nome
        caminho.write_bytes(conteudo)
        caminhos.append(caminho)
    return caminhos


if __name__ == "__main__":
    for c in gerar_todos(RAIZ / "exemplos"):
        print("gerado:", c.relative_to(RAIZ))
    sys.exit(0)
