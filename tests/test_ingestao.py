from __future__ import annotations

import io
import zipfile
from datetime import datetime

import openpyxl
import pytest

import gerar_exemplos as fab
from schulz import ingestao as ing
from schulz.ingestao import (ErroPlanilha, data_iso, ler_planilha, numero,
                                         precisa_reparo, repair_xlsx_stream)


def test_modelo_simples_aprovado(xlsx_aprovado):
    pl = ler_planilha(xlsx_aprovado, "a.xlsx")
    assert pl.formato == "tabela_de_leituras" and pl.aba == "Leituras"
    assert len(pl.pontos) == 10
    p0 = pl.pontos[0]
    assert (p0.nome, p0.valor_padrao, p0.lea_positivo, p0.lea_negativo) == ("DIAMETRO_A", 10.0, 0.01, 0.01)
    assert p0.leituras == (10.0021, 10.0027, 10.0024)


def test_metadados_extraidos_por_palavra_chave():
    meta = {"Cliente:": "Usinagem Teste", "Técnico:": "Fulano", "Código:": "123.456",
            "Denominação:": "Calibrador", "Data:": datetime(2026, 9, 1)}
    pl = ler_planilha(fab.tabela_simples(fab.PONTOS_APROVADO, meta=meta), "a.xlsx")
    m = pl.metadados
    assert (m["cliente"], m["tecnico"], m["codigo_instrumento"], m["denominacao"]) == \
        ("Usinagem Teste", "Fulano", "123.456", "Calibrador")
    assert m["data_calibracao"] == "2026-09-01"


def test_cabecalho_com_sinonimos_e_linhas_de_texto_no_topo():
    dados = fab.tabela_simples(
        [("P1", 5.0, 0.02, 5.001, 5.003), ("P2", 7.0, 0.02, 6.999, 7.001)],
        cabecalho=["Parâmetro", "Nominal", "LEA", "Medida 1", "Medida 2"],
        meta={"Relatório interno": "texto livre", "Observação": "ignorar"}, linhas_topo=3)
    pl = ler_planilha(dados, "s.xlsx")
    assert [p.nome for p in pl.pontos] == ["P1", "P2"]


def test_virgula_decimal_em_texto():
    dados = fab.tabela_simples([("P1", "6,075", "0,010", "6,0781", "6,0779")])
    p = ler_planilha(dados, "v.xlsx").pontos[0]
    assert p.valor_padrao == 6.075 and p.lea_positivo == 0.01 and p.leituras == (6.0781, 6.0779)


def test_quantidade_variavel_de_leituras():
    dados = fab.tabela_simples([("P1", 1.0, 0.1, 1.0, 1.02, 0.98, 1.0)])
    assert len(ler_planilha(dados, "v.xlsx").pontos[0].leituras) == 4


def test_rodape_notas_e_residuos_numericos_sao_ignorados():
    base = [("P1", 5.0, 0.02, 5.001, 5.003), ("P2", 7.0, 0.02, 6.999, 7.001),
            ("Nota: incerteza expandida k=2", None, None, None, None),
            ("2", 0, 0, 0, 0), ("0", 1, 1, 1, 1), ("1", 2, 2, 2, 2),
            ("*Observação técnica", 1, 1, 1, 1)]
    pl = ler_planilha(fab.tabela_simples(base), "r.xlsx")
    assert [p.nome for p in pl.pontos] == ["P1", "P2"]


def test_abas_extras_sao_ignoradas_e_a_prioridade_vence():
    wb = openpyxl.Workbook()
    capa = wb.active
    capa.title = "Capa"
    capa["A1"] = "Documento sem tabela"
    ws = wb.create_sheet("Registro")
    for c, t in enumerate(["Ponto", "Valor padrão", "LEA", "Leitura 1", "Leitura 2"], 1):
        ws.cell(row=2, column=c, value=t)
    for c, v in enumerate(["X", 1.0, 0.1, 1.01, 0.99], 1):
        ws.cell(row=3, column=c, value=v)
    buf = io.BytesIO()
    wb.save(buf)
    pl = ler_planilha(buf.getvalue(), "m.xlsx")
    assert pl.aba == "Registro" and pl.pontos[0].nome == "X"


def test_tolerancias_assimetricas_preservadas():
    dados = fab.tabela_tolerancias([("PLANICIDADE", 0.0, 0.05, 0.0, 0.002, 0.003),
                                    ("DIAMETRO", 6.075, 0.01, -0.01, 6.078, 6.079)])
    a, b = ler_planilha(dados, "t.xlsx").pontos
    assert (a.lea_positivo, a.lea_negativo) == (0.05, 0.0)
    assert (b.lea_positivo, b.lea_negativo) == (0.01, 0.01)


def test_calypso_blocos_viram_leituras_e_metadados_sao_lidos():
    dados = fab.calypso([fab.BLOCO_CALYPSO_1, fab.BLOCO_CALYPSO_2], plano="703.1208_TESTE",
                        metrologista="Fulano de Tal", mmc=500008, fornecedor="CDM",
                        data=datetime(2026, 8, 13))
    pl = ler_planilha(dados, "calypso.xlsx")
    assert pl.formato == "calypso_report" and pl.aba == "Report"
    assert len(pl.pontos) == 6
    p = pl.pontos[0]
    assert p.nome == "PLANICIDADE_REF" and p.leituras == (0.0021, 0.0023)
    assert (p.valor_padrao, p.lea_positivo, p.lea_negativo) == (0.0, 0.05, 0.0)
    m = pl.metadados
    assert m["tecnico"] == "Fulano de Tal" and m["plano_medicao"] == "703.1208_TESTE"
    assert m["padrao_utilizado"] == "500008" and m["fabricante"] == "CDM"
    assert m["data_calibracao"] == "2026-08-13" and m["codigo_instrumento"] == "703.1208"


def test_calypso_rodape_nao_vira_ponto():
    pl = ler_planilha(fab.calypso([fab.BLOCO_CALYPSO_1, fab.BLOCO_CALYPSO_2], rodape=True), "c.xlsx")
    assert all(p.nome[0].isalpha() and not p.nome.startswith("Nota") for p in pl.pontos)


def test_calypso_tres_blocos_geram_tres_leituras():
    b3 = [(n, a + 0.0001, nom, up, lo) for n, a, nom, up, lo in fab.BLOCO_CALYPSO_1]
    pl = ler_planilha(fab.calypso([fab.BLOCO_CALYPSO_1, fab.BLOCO_CALYPSO_2, b3]), "c.xlsx")
    assert all(len(p.leituras) == 3 for p in pl.pontos)


def test_calypso_com_um_unico_bloco_e_rejeitado():
    with pytest.raises(ErroPlanilha) as e:
        ler_planilha(fab.calypso([fab.BLOCO_CALYPSO_1]), "c.xlsx")
    assert e.value.codigo == "LEITURAS_INSUFICIENTES" and e.value.status == 422


def test_calypso_blocos_inconsistentes_sao_rejeitados():
    incompleto = fab.BLOCO_CALYPSO_2[:-1]
    with pytest.raises(ErroPlanilha) as e:
        ler_planilha(fab.calypso([fab.BLOCO_CALYPSO_1, incompleto]), "c.xlsx")
    assert e.value.codigo == "DADOS_INCOMPLETOS"
    assert any("ENTRE_CENTROS" in d for d in e.value.detalhes)


def test_registro_schulz_em_formato_estrito():
    original = fab.registro_schulz(fab.PONTOS_REGISTRO)
    estrito = fab.para_estrito(original)
    assert precisa_reparo(estrito) and not precisa_reparo(original)
    assert openpyxl.load_workbook(io.BytesIO(estrito)).sheetnames == []
    pl = ler_planilha(estrito, "registro.xlsx")
    assert pl.aba == "Registro" and len(pl.pontos) == 5
    assert pl.avisos[0].startswith("Arquivo em OOXML estrito")
    assert pl.metadados["cliente"] == "Setor Exemplo" and pl.metadados["tecnico"] == "Técnico Exemplo"
    assert pl.metadados["codigo_instrumento"] == "000.0000"
    assert pl.metadados["data_calibracao"] == "2026-09-01"
    assert pl.metadados["temperatura_ambiente"] == "20,2°"


def test_colunas_calculadas_e_erros_de_formula_sao_ignorados():
    pl = ler_planilha(fab.registro_schulz(fab.PONTOS_REGISTRO), "registro.xlsx")
    assert len(pl.pontos) == 5
    assert pl.pontos[1].leituras == (6.0781, 6.0779)
    assert pl.pontos[2].lea_negativo == 0.0


def test_repair_xlsx_stream_devolve_bytesio_com_namespaces_universais():
    estrito = fab.para_estrito(fab.tabela_simples(fab.PONTOS_APROVADO))
    saida = repair_xlsx_stream(estrito)
    assert isinstance(saida, io.BytesIO)
    with zipfile.ZipFile(saida) as z:
        wbxml = z.read("xl/workbook.xml").decode()
    assert "purl.oclc.org" not in wbxml and "schemas.openxmlformats.org/spreadsheetml/2006/main" in wbxml
    assert 'conformance="strict"' not in wbxml
    saida.seek(0)
    assert openpyxl.load_workbook(saida).sheetnames == ["Leituras"]


def test_reparo_e_idempotente_em_arquivo_normal(xlsx_aprovado):
    r1 = repair_xlsx_stream(xlsx_aprovado).getvalue()
    assert openpyxl.load_workbook(io.BytesIO(r1)).sheetnames == ["Leituras"]
    assert len(ler_planilha(r1, "x.xlsx").pontos) == 10


def test_arquivo_que_nao_e_zip():
    with pytest.raises(ErroPlanilha) as e:
        ler_planilha(b"isto nao e um excel", "x.xlsx")
    assert e.value.codigo == "ARQUIVO_INVALIDO" and e.value.status == 400


@pytest.mark.parametrize("nome", ["dados.csv", "dados.xls", "dados.pdf", "", "semextensao"])
def test_extensao_nao_suportada(nome):
    with pytest.raises(ErroPlanilha) as e:
        ler_planilha(b"x", nome)
    assert e.value.codigo == "ARQUIVO_INVALIDO" and e.value.status == 400


def test_planilha_sem_cabecalho_reconhecivel():
    wb = openpyxl.Workbook()
    wb.active["A1"] = "qualquer coisa"
    wb.active["B2"] = 42
    buf = io.BytesIO()
    wb.save(buf)
    with pytest.raises(ErroPlanilha) as e:
        ler_planilha(buf.getvalue(), "x.xlsx")
    assert e.value.codigo == "LAYOUT_NAO_RECONHECIDO" and "Abas examinadas" in e.value.detalhes[0]


def test_apenas_uma_coluna_de_leitura():
    dados = fab.tabela_simples([("P1", 1.0, 0.1, 1.0)], cabecalho=["Ponto", "Valor padrão", "LEA", "Leitura 1"])
    with pytest.raises(ErroPlanilha) as e:
        ler_planilha(dados, "x.xlsx")
    assert e.value.codigo == "LEITURAS_INSUFICIENTES"


def test_pontos_incompletos_bloqueiam_e_indicam_a_linha():
    dados = fab.tabela_simples([("OK", 1.0, 0.1, 1.0, 1.0), ("SEM_PADRAO", None, 0.1, 1.0, 1.0),
                                ("SEM_LEA", 2.0, None, 2.0, 2.0), ("UMA_LEITURA", 3.0, 0.1, 3.0, None)])
    with pytest.raises(ErroPlanilha) as e:
        ler_planilha(dados, "x.xlsx")
    assert e.value.codigo == "DADOS_INCOMPLETOS" and len(e.value.detalhes) == 3
    assert "Linha 3" in e.value.detalhes[0] and "valor padrão" in e.value.detalhes[0]
    assert "tolerância/LEA" in e.value.detalhes[1] and "leituras" in e.value.detalhes[2]


def test_planilha_somente_com_cabecalho():
    with pytest.raises(ErroPlanilha) as e:
        ler_planilha(fab.tabela_simples([]), "x.xlsx")
    assert e.value.codigo in {"SEM_PONTOS", "LAYOUT_NAO_RECONHECIDO"}


@pytest.mark.parametrize("bruto, esperado", [
    (1, 1.0), (2.5, 2.5), ("0,5", 0.5), ("6.075", 6.075), ("1.234,56", 1234.56), (" 3,0 ", 3.0),
    ("#VALUE!", None), ("#DIV/0!", None), ("----", None), ("???", None), (" ", None), ("", None),
    (None, None), (True, None), (float("inf"), None), (float("nan"), None), ("abc", None),
])
def test_numero(bruto, esperado):
    assert numero(bruto) == esperado


@pytest.mark.parametrize("bruto, esperado", [
    (datetime(2026, 8, 13, 16, 16), "2026-08-13"), (46000, "2025-12-09"), ("13/08/2026", "2026-08-13"),
    ("2026-08-13", "2026-08-13"), ("00/00/0000", None), ("lixo", None), (None, None), (12, None),
])
def test_data_iso(bruto, esperado):
    assert data_iso(bruto) == esperado


def test_normalizar_cabecalhos():
    n = ing.normalizar
    assert n("Tol. +") == n("tol+") == n("TOL +") == "tol +"
    assert n("+TOL") == "+tol" and n("Tol. −") == "tol -" and n("Valor padrão (mm)") == "valor padrao"
    assert n("Parâmetro avaliado") == "parametro avaliado"
