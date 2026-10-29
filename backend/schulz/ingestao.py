from __future__ import annotations

import io
import math
import re
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime, time, timedelta
from typing import Any, BinaryIO

import openpyxl

from .calculo import Ponto

LIMITE_DESCOMPACTADO = 100 * 1024 * 1024
LIMITE_LINHAS = 5000
LIMITE_COLUNAS = 80
EXTENSOES_ACEITAS = (".xlsx",)

_NAMESPACES_ESTRITO_TRANSICIONAL = (
    ("http://purl.oclc.org/ooxml/spreadsheetml/main",
     "http://schemas.openxmlformats.org/spreadsheetml/2006/main"),
    ("http://purl.oclc.org/ooxml/officeDocument/relationships",
     "http://schemas.openxmlformats.org/officeDocument/2006/relationships"),
    ("http://purl.oclc.org/ooxml/drawingml/main",
     "http://schemas.openxmlformats.org/drawingml/2006/main"),
    ("http://purl.oclc.org/ooxml/drawingml/spreadsheetDrawing",
     "http://schemas.openxmlformats.org/drawingml/2006/spreadsheetDrawing"),
    ("http://purl.oclc.org/ooxml/drawingml/chart",
     "http://schemas.openxmlformats.org/drawingml/2006/chart"),
    ("http://purl.oclc.org/ooxml/officeDocument/sharedTypes",
     "http://schemas.openxmlformats.org/officeDocument/2006/sharedTypes"),
    ("http://purl.oclc.org/ooxml/officeDocument/docPropsVTypes",
     "http://schemas.openxmlformats.org/officeDocument/2006/docPropsVTypes"),
    ("http://purl.oclc.org/ooxml/officeDocument/extendedProperties",
     "http://schemas.openxmlformats.org/officeDocument/2006/extended-properties"),
)


class ErroPlanilha(Exception):
    def __init__(self, codigo: str, mensagem: str, detalhes: list[str] | None = None,
                 status: int = 422):
        super().__init__(mensagem)
        self.codigo = codigo
        self.mensagem = mensagem
        self.detalhes = detalhes or []
        self.status = status

    def como_dict(self) -> dict[str, Any]:
        return {"codigo": self.codigo, "mensagem": self.mensagem, "detalhes": self.detalhes}


@dataclass
class PlanilhaLida:
    formato: str
    aba: str
    metadados: dict[str, str | None]
    pontos: list[Ponto]
    avisos: list[str] = field(default_factory=list)


def precisa_reparo(dados: bytes) -> bool:
    try:
        with zipfile.ZipFile(io.BytesIO(dados)) as z:
            return b"purl.oclc.org/ooxml" in z.read("xl/workbook.xml")
    except (zipfile.BadZipFile, KeyError):
        return False


def reparar_xlsx_stream(origem: bytes | str | BinaryIO) -> io.BytesIO:
    if isinstance(origem, (bytes, bytearray)):
        entrada: Any = io.BytesIO(bytes(origem))
    else:
        entrada = origem
    saida = io.BytesIO()
    with zipfile.ZipFile(entrada, "r") as z_in, \
            zipfile.ZipFile(saida, "w", zipfile.ZIP_DEFLATED) as z_out:
        if sum(i.file_size for i in z_in.infolist()) > LIMITE_DESCOMPACTADO:
            raise ErroPlanilha("ARQUIVO_INVALIDO", "Arquivo descompactado excede o limite permitido.",
                               status=400)
        for item in z_in.infolist():
            dados = z_in.read(item.filename)
            if item.filename.endswith((".xml", ".rels")):
                texto = dados.decode("utf-8")
                for estrito, transicional in _NAMESPACES_ESTRITO_TRANSICIONAL:
                    texto = texto.replace(estrito, transicional)
                texto = texto.replace(' conformance="strict"', "")
                dados = texto.encode("utf-8")
            z_out.writestr(item, dados)
    saida.seek(0)
    return saida


repair_xlsx_stream = reparar_xlsx_stream


def carregar_workbook(dados: bytes) -> tuple[openpyxl.Workbook, bool]:
    if not zipfile.is_zipfile(io.BytesIO(dados)):
        raise ErroPlanilha(
            "ARQUIVO_INVALIDO",
            "O arquivo não é um Excel .xlsx válido (conteúdo não reconhecido).",
            status=400,
        )
    reparado = precisa_reparo(dados)
    try:
        fonte = reparar_xlsx_stream(dados) if reparado else io.BytesIO(dados)
        wb = openpyxl.load_workbook(fonte, data_only=True)
    except ErroPlanilha:
        raise
    except Exception:
        try:
            wb = openpyxl.load_workbook(reparar_xlsx_stream(dados), data_only=True)
            reparado = True
        except Exception as exc:
            raise ErroPlanilha(
                "ARQUIVO_INVALIDO",
                "Não foi possível abrir o arquivo Excel (corrompido ou protegido por senha).",
                [type(exc).__name__],
                status=400,
            ) from exc
    if not wb.worksheets:
        raise ErroPlanilha("ARQUIVO_INVALIDO", "A planilha não contém nenhuma aba legível.",
                           status=400)
    return wb, reparado


def normalizar(texto: Any) -> str:
    if texto is None:
        return ""
    s = str(texto).strip().lower()
    s = s.replace("−", "-").replace("–", "-").replace("—", "-")
    s = "".join(c for c in unicodedata.normalize("NFKD", s) if not unicodedata.combining(c))
    s = re.sub(r"\([^)]*\)", " ", s)
    s = s.replace(".", " ").replace(":", " ").replace("*", " ")
    s = re.sub(r"\s+", " ", s).strip()
    s = re.sub(r"\b(tol|lea)\s*([+-])", r"\1 \2", s)
    s = re.sub(r"([+-])\s*(tol|lea)\b", r"\1\2", s)
    return s


_PLACEHOLDERS = {
    "", "-", "--", "---", "----", "-----", "------", "???", "xx", "yy", "n/a", "na", "nd",
    "#value!", "#div/0!", "#ref!", "#n/a", "#name?", "#num!", "#null!", "00/00/0000",
}


def texto_limpo(valor: Any) -> str | None:
    if valor is None:
        return None
    if isinstance(valor, (datetime, date, time)):
        return None
    s = str(valor).strip()
    if s.lower() in _PLACEHOLDERS:
        return None
    return s


def numero(valor: Any) -> float | None:
    if valor is None or isinstance(valor, bool):
        return None
    if isinstance(valor, (int, float)):
        return float(valor) if math.isfinite(valor) else None
    if isinstance(valor, (datetime, date, time)):
        return None
    s = str(valor).strip()
    if s.lower() in _PLACEHOLDERS:
        return None
    s = s.replace("−", "-").replace(" ", "")
    if "," in s and "." in s:
        s = s.replace(".", "").replace(",", ".")
    else:
        s = s.replace(",", ".")
    try:
        v = float(s)
    except ValueError:
        return None
    return v if math.isfinite(v) else None


def data_iso(valor: Any) -> str | None:
    if isinstance(valor, datetime):
        return valor.date().isoformat()
    if isinstance(valor, date):
        return valor.isoformat()
    if isinstance(valor, (int, float)) and not isinstance(valor, bool):
        if 20000 <= valor <= 80000:
            return (datetime(1899, 12, 30) + timedelta(days=float(valor))).date().isoformat()
        return None
    t = texto_limpo(valor)
    if not t:
        return None
    for fmt in ("%d/%m/%Y", "%Y-%m-%d", "%d/%m/%y", "%d-%m-%Y"):
        try:
            return datetime.strptime(t.split(" ")[0], fmt).date().isoformat()
        except ValueError:
            continue
    return None


def _tem_letra(texto: str) -> bool:
    return any(c.isalpha() for c in texto)


_PREFIXOS_NOTA = ("*", "nota", "obs", "observ", "incerteza", "fonte", "legenda")


def _nome_de_ponto_valido(nome: str | None) -> bool:
    if not nome or not _tem_letra(nome):
        return False
    n = normalizar(nome)
    if not n or n.startswith(_PREFIXOS_NOTA) or nome.strip().startswith("*"):
        return False
    return n not in {"media", "total", "resultado", "parecer", "desvio padrao"}


Grade = list[list[Any]]


def _grade(ws) -> Grade:
    linhas: Grade = []
    for i, row in enumerate(ws.iter_rows(max_col=LIMITE_COLUNAS, values_only=True)):
        if i >= LIMITE_LINHAS:
            break
        linhas.append(list(row))
    while linhas and all(c is None or str(c).strip() == "" for c in linhas[-1]):
        linhas.pop()
    return linhas


def _linha_vazia(linha: list[Any]) -> bool:
    return all(c is None or (isinstance(c, str) and c.strip() == "") for c in linha)


_PAPEIS_TABELA: dict[str, set[str]] = {
    "ponto": {"parametro avaliado", "parametro", "dimension", "dimensao", "ponto",
              "ponto de ensaio", "caracteristica", "characteristic", "ponto de medicao"},
    "padrao": {"valor nominal", "nominal", "padrao", "valor padrao", "valor do padrao",
               "valor de referencia", "referencia", "valor de referencia do padrao"},
    "tol_pos": {"tol +", "+tol", "tolerancia +", "tolerancia superior", "upper tol",
                "lea +", "lea superior", "tol sup"},
    "tol_neg": {"tol -", "-tol", "tolerancia -", "tolerancia inferior", "lower tol",
                "lea -", "lea inferior", "tol inf"},
    "lea": {"lea", "limite de erro admissivel", "erro admissivel", "tolerancia"},
}
_RE_LEITURA = re.compile(r"^(medida|meas|leitura|medicao|med|l|m)\s*(\d{1,2})$")

_PAPEIS_CALYPSO: dict[str, set[str]] = {
    "ponto": {"characteristic", "caracteristica"},
    "actual": {"actual", "valor medido", "medido", "obtido"},
    "padrao": {"nominal", "valor nominal"},
    "tol_pos": {"upper tol", "tol +", "+tol"},
    "tol_neg": {"lower tol", "tol -", "-tol"},
}


def _mapear_cabecalho(linha: list[Any], papeis: dict[str, set[str]],
                      com_leituras: bool) -> dict[str, Any]:
    mapa: dict[str, Any] = {}
    leituras: dict[int, int] = {}
    for col, celula in enumerate(linha):
        rot = normalizar(celula)
        if not rot:
            continue
        if com_leituras:
            m = _RE_LEITURA.match(rot)
            if m:
                leituras.setdefault(int(m.group(2)), col)
                continue
        for papel, nomes in papeis.items():
            if rot in nomes and papel not in mapa:
                mapa[papel] = col
                break
    if com_leituras and leituras:
        mapa["leituras"] = [leituras[k] for k in sorted(leituras)]
    return mapa


def _celula(linha: list[Any], col: int | None) -> Any:
    if col is None or col >= len(linha):
        return None
    return linha[col]


_ROTULOS_META: dict[str, str] = {
    "cliente solicitante": "cliente", "cliente": "cliente", "solicitante": "cliente",
    "tecnico executante": "tecnico", "tecnico": "tecnico", "metrologista": "tecnico",
    "operador": "tecnico",
    "codigo": "codigo_instrumento", "codigo do instrumento": "codigo_instrumento",
    "tag": "codigo_instrumento",
    "denominacao": "denominacao", "instrumento": "denominacao",
    "descricao do instrumento": "denominacao",
    "fabricante": "fabricante", "fornecedor": "fabricante",
    "data": "data_calibracao", "date": "data_calibracao", "data da calibracao": "data_calibracao",
    "temperatura ambiente": "temperatura_ambiente", "temperatura": "temperatura_ambiente",
    "measurement plan": "plano_medicao", "part program name": "plano_medicao",
    "mmc": "padrao_utilizado",
}
_CAMPOS_META = sorted(set(_ROTULOS_META.values()))
_OUTROS_ROTULOS = {"order", "drawing no", "time", "part no", "part number", "ordem", "hora"}


def _metadados_vazios() -> dict[str, str | None]:
    return {c: None for c in _CAMPOS_META}


def _valor_meta(campo: str, bruto: Any) -> str | None:
    if campo == "data_calibracao":
        return data_iso(bruto)
    if isinstance(bruto, float) and bruto.is_integer():
        bruto = int(bruto)
    return texto_limpo(bruto)


def _metadados(grade: Grade, ate_linha: int) -> dict[str, str | None]:
    meta = _metadados_vazios()
    rotulos = set(_ROTULOS_META) | _OUTROS_ROTULOS
    for r in range(min(ate_linha, len(grade))):
        linha = grade[r]
        for c, cel in enumerate(linha):
            if not isinstance(cel, str):
                continue
            rot = normalizar(cel)
            campo = _ROTULOS_META.get(rot)
            if campo is None or meta[campo] is not None:
                continue
            candidato: Any = None
            for cc in range(c + 1, min(c + 6, len(linha))):
                v = linha[cc]
                if v is None or (isinstance(v, str) and not v.strip()):
                    continue
                if isinstance(v, str) and (v.strip().endswith(":") or normalizar(v) in rotulos):
                    break
                candidato = v
                break
            if candidato is None and r + 1 < len(grade):
                abaixo = grade[r + 1][c] if c < len(grade[r + 1]) else None
                if abaixo is not None and not (isinstance(abaixo, str) and
                                                (not abaixo.strip() or normalizar(abaixo) in rotulos)):
                    candidato = abaixo
            valor = _valor_meta(campo, candidato)
            if valor is not None:
                meta[campo] = valor
    return meta


def _ler_tabela(grade: Grade, aba: str) -> PlanilhaLida | None:
    cabecalho = None
    for idx, linha in enumerate(grade[:300]):
        mapa = _mapear_cabecalho(linha, _PAPEIS_TABELA, com_leituras=True)
        tem_tol = ("tol_pos" in mapa and "tol_neg" in mapa) or "lea" in mapa
        if "ponto" in mapa and "padrao" in mapa and tem_tol and len(mapa.get("leituras", [])) >= 1:
            cabecalho = (idx, mapa)
            break
    if cabecalho is None:
        return None
    idx, mapa = cabecalho
    colunas_leitura: list[int] = mapa["leituras"]
    if len(colunas_leitura) < 2:
        raise ErroPlanilha(
            "LEITURAS_INSUFICIENTES",
            "O cabeçalho tem apenas uma coluna de leitura; são necessárias ao menos duas "
            "(ex.: 'Medida 1' e 'Medida 2') para média e desvio padrão.",
        )
    usar_tol = "tol_pos" in mapa and "tol_neg" in mapa

    pontos: list[Ponto] = []
    incompletos: list[str] = []
    avisos: list[str] = []
    vazias_seguidas = 0
    for r in range(idx + 1, len(grade)):
        linha = grade[r]
        nome = texto_limpo(_celula(linha, mapa["ponto"]))
        numero_linha = r + 1
        if not nome:
            vazias_seguidas += 1
            if pontos and vazias_seguidas >= 5:
                break
            continue
        vazias_seguidas = 0
        if not _nome_de_ponto_valido(nome):
            continue
        padrao = numero(_celula(linha, mapa["padrao"]))
        leituras = [v for v in (numero(_celula(linha, c)) for c in colunas_leitura) if v is not None]
        if usar_tol:
            tp, tn = numero(_celula(linha, mapa["tol_pos"])), numero(_celula(linha, mapa["tol_neg"]))
        else:
            lea = numero(_celula(linha, mapa["lea"]))
            tp = tn = lea
        if padrao is None and not leituras and tp is None and tn is None:
            if len(avisos) < 10:
                avisos.append(f"Linha {numero_linha} ('{nome}') ignorada: sem valores numéricos.")
            continue
        faltas = []
        if padrao is None:
            faltas.append("valor padrão")
        if tp is None or tn is None:
            faltas.append("tolerância/LEA")
        if len(leituras) < 2:
            faltas.append(f"leituras (há {len(leituras)}, mínimo 2)")
        if faltas:
            incompletos.append(f"Linha {numero_linha} – '{nome}': faltando {', '.join(faltas)}.")
            continue
        if len(leituras) < len(colunas_leitura):
            avisos.append(f"Linha {numero_linha} – '{nome}': {len(colunas_leitura) - len(leituras)} "
                          f"leitura(s) em branco foram desconsideradas.")
        pontos.append(Ponto(nome=nome, valor_padrao=padrao, lea_positivo=abs(tp),
                            lea_negativo=abs(tn), leituras=tuple(leituras)))
    if incompletos:
        raise ErroPlanilha(
            "DADOS_INCOMPLETOS",
            f"{len(incompletos)} ponto(s) com dados incompletos; corrija a planilha e envie novamente.",
            incompletos,
        )
    if not pontos:
        raise ErroPlanilha("SEM_PONTOS", "O cabeçalho foi encontrado, mas não há pontos de ensaio preenchidos abaixo dele (planilha em branco?).")
    return PlanilhaLida("tabela_de_leituras", aba, _metadados(grade, idx), pontos, avisos)


def _ler_calypso(grade: Grade, aba: str) -> PlanilhaLida | None:
    cabecalho = None
    for idx, linha in enumerate(grade[:300]):
        mapa = _mapear_cabecalho(linha, _PAPEIS_CALYPSO, com_leituras=False)
        if "ponto" in mapa and "actual" in mapa:
            cabecalho = (idx, mapa)
            break
    if cabecalho is None:
        return None
    idx, mapa = cabecalho
    blocos: list[dict[str, tuple[float, float | None, float | None, float | None]]] = []
    atual: dict[str, tuple] = {}
    avisos: list[str] = []
    incompletos: list[str] = []
    ordem: list[str] = []

    def fechar() -> None:
        nonlocal atual
        if atual:
            blocos.append(atual)
            atual = {}

    for r in range(idx + 1, len(grade)):
        linha = grade[r]
        if _linha_vazia(linha):
            fechar()
            continue
        nome = texto_limpo(_celula(linha, mapa["ponto"]))
        if not _nome_de_ponto_valido(nome):
            continue
        actual = numero(_celula(linha, mapa["actual"]))
        if actual is None:
            incompletos.append(f"Linha {r + 1} – '{nome}': valor 'Actual' ausente ou inválido.")
            continue
        if nome in atual:
            fechar()
        if nome not in ordem:
            ordem.append(nome)
        atual[nome] = (
            actual,
            numero(_celula(linha, mapa.get("padrao"))),
            numero(_celula(linha, mapa.get("tol_pos"))),
            numero(_celula(linha, mapa.get("tol_neg"))),
        )
    fechar()
    if incompletos:
        raise ErroPlanilha("DADOS_INCOMPLETOS", f"{len(incompletos)} linha(s) com valor ausente.", incompletos)
    if not blocos:
        raise ErroPlanilha("SEM_PONTOS", "O cabeçalho do Calypso foi encontrado, mas não há medições preenchidas abaixo dele (relatório em branco?).")
    if len(blocos) < 2:
        raise ErroPlanilha(
            "LEITURAS_INSUFICIENTES",
            "Foi encontrado apenas 1 bloco de medição no relatório do Calypso; a calibração "
            "exige duas rodadas de leituras por característica (RN05).",
        )
    pontos: list[Ponto] = []
    for nome in blocos[0]:
        leituras, divergencias = [], []
        for n_bloco, bloco in enumerate(blocos, start=1):
            if nome not in bloco:
                incompletos.append(f"'{nome}' não aparece no bloco {n_bloco} do relatório.")
                break
            leituras.append(bloco[nome][0])
            if bloco[nome][1:] != blocos[0][nome][1:]:
                divergencias.append(n_bloco)
        else:
            _, nominal, tp, tn = blocos[0][nome]
            if nominal is None or tp is None or tn is None:
                incompletos.append(f"'{nome}': Nominal/Upper Tol/Lower Tol ausentes.")
                continue
            if divergencias:
                avisos.append(f"'{nome}': nominal/tolerância diferem nos blocos {divergencias}; "
                              f"foi usado o bloco 1.")
            pontos.append(Ponto(nome=nome, valor_padrao=nominal, lea_positivo=abs(tp),
                                lea_negativo=abs(tn), leituras=tuple(leituras)))
    for n_bloco, bloco in enumerate(blocos[1:], start=2):
        extras = [n for n in bloco if n not in blocos[0]]
        if extras:
            incompletos.append(f"Bloco {n_bloco} tem característica(s) ausente(s) no bloco 1: {', '.join(extras)}.")
    if incompletos:
        raise ErroPlanilha("DADOS_INCOMPLETOS", "Os blocos do relatório não são consistentes entre si.", incompletos)
    meta = _metadados(grade, idx)
    if meta["codigo_instrumento"] is None and meta["plano_medicao"]:
        achado = re.search(r"(?<!\d)\d{3}\.\d{3,5}(?!\d)", meta["plano_medicao"])
        if achado:
            meta["codigo_instrumento"] = achado.group(0)
            avisos.append("Código do instrumento inferido do nome do plano de medição.")
    avisos.append(f"Relatório Calypso com {len(blocos)} bloco(s) de medição: cada bloco foi tratado como uma leitura.")
    return PlanilhaLida("calypso_report", aba, meta, pontos, avisos)


_PRIORIDADE_ABAS = ("registro", "importacao", "report", "dados", "leituras", "medicoes", "calibracao")


def _ordenar_abas(nomes: list[str]) -> list[str]:
    def chave(par: tuple[int, str]) -> tuple[int, int]:
        pos, nome = par
        n = normalizar(nome)
        return (_PRIORIDADE_ABAS.index(n) if n in _PRIORIDADE_ABAS else len(_PRIORIDADE_ABAS), pos)
    return [n for _, n in sorted(enumerate(nomes), key=chave)]


def validar_nome_arquivo(nome: str | None) -> str:
    nome = (nome or "").strip()
    if not nome.lower().endswith(EXTENSOES_ACEITAS):
        raise ErroPlanilha(
            "ARQUIVO_INVALIDO",
            "Formato não suportado. Envie um arquivo Excel .xlsx (arquivos com macros .xlsm são recusados por segurança).",
            [f"Arquivo recebido: '{nome or 'sem nome'}'"],
            status=400,
        )
    return nome


def ler_planilha(dados: bytes, nome_arquivo: str = "planilha.xlsx") -> PlanilhaLida:
    validar_nome_arquivo(nome_arquivo)
    wb, reparado = carregar_workbook(dados)
    primeiro_erro: ErroPlanilha | None = None
    for nome_aba in _ordenar_abas([ws.title for ws in wb.worksheets]):
        grade = _grade(wb[nome_aba])
        if not grade:
            continue
        for leitor in (_ler_tabela, _ler_calypso):
            try:
                resultado = leitor(grade, nome_aba)
            except ErroPlanilha as exc:
                primeiro_erro = primeiro_erro or exc
                continue
            if resultado is not None:
                if reparado:
                    resultado.avisos.insert(0, "Arquivo em OOXML estrito: namespaces saneados em memória antes da leitura.")
                return resultado
    if primeiro_erro is not None:
        raise primeiro_erro
    raise ErroPlanilha(
        "LAYOUT_NAO_RECONHECIDO",
        "Não encontrei uma tabela de leituras. O cabeçalho deve conter, por exemplo: "
        "'Ponto' (ou 'Parâmetro avaliado'), 'Valor padrão' (ou 'Valor nominal'), "
        "'LEA' (ou 'Tol. +' e 'Tol. -') e 'Leitura 1', 'Leitura 2'. "
        "Também é aceito o relatório do Calypso (Characteristic / Actual / Nominal / Upper Tol / Lower Tol).",
        [f"Abas examinadas: {', '.join(ws.title for ws in wb.worksheets)}"],
    )
